import os, re, io, csv, uuid, secrets, zipfile, threading, hashlib

from pathlib import Path, PurePosixPath

from datetime import datetime, timezone

from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException

from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse, FileResponse

from fastapi.staticfiles import StaticFiles

from fastapi.templating import Jinja2Templates

from starlette.middleware.sessions import SessionMiddleware

from authlib.integrations.starlette_client import OAuth

from sqlalchemy import create_engine, text

import pymupdf

from docx import Document

from sklearn.feature_extraction.text import TfidfVectorizer

from sklearn.metrics.pairwise import cosine_similarity

from openpyxl import Workbook

from openpyxl.styles import Font, Alignment, PatternFill



BASE=Path(__file__).resolve().parent.parent

URL=os.getenv('DATABASE_URL',f"sqlite:///{(BASE/'resumeflow.db').as_posix()}")

if URL.startswith("postgres://"):
    URL = URL.replace(
        "postgres://",
        "postgresql+psycopg://",
        1,
    )
elif URL.startswith("postgresql://"):
    URL = URL.replace(
        "postgresql://",
        "postgresql+psycopg://",
        1,
    )

kw={'pool_pre_ping':True}

if URL.startswith('sqlite'): kw['connect_args']={'check_same_thread':False}

engine=create_engine(URL,**kw)

SECRET=os.getenv('SECRET_KEY') or secrets.token_urlsafe(48)

MAX_MB=int(os.getenv('MAX_UPLOAD_MB','25')); MAX_FILES=int(os.getenv('MAX_FILES_PER_BATCH','500')); MAX_ZIP_MB=int(os.getenv('MAX_ZIP_UNCOMPRESSED_MB','300'))

ALLOWED={x.strip().lower() for x in os.getenv('ALLOWED_EMAILS','').split(',') if x.strip()}

GOOGLE=bool(os.getenv('GOOGLE_CLIENT_ID') and os.getenv('GOOGLE_CLIENT_SECRET'))

DEMO_LOGIN_ENABLED = (
    os.getenv("DEMO_LOGIN_ENABLED", "").lower() == "true"
)

app=FastAPI(title='ResumeFlow | AI Recruitment Workspace')

app.add_middleware(SessionMiddleware,secret_key=SECRET,https_only=os.getenv('APP_ENV')=='production',same_site='lax')

app.mount('/static',StaticFiles(directory=str(BASE/'app'/'static')),name='static')

templates=Jinja2Templates(directory=str(BASE/'app'/'templates'))

oauth=OAuth()

if GOOGLE: oauth.register(name='google',client_id=os.getenv('GOOGLE_CLIENT_ID'),client_secret=os.getenv('GOOGLE_CLIENT_SECRET'),server_metadata_url='https://accounts.google.com/.well-known/openid-configuration',client_kwargs={'scope':'openid email profile'})

SCHEMA = """
CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY,email TEXT NOT NULL UNIQUE,name TEXT NOT NULL,picture TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,user_id TEXT NOT NULL,title TEXT NOT NULL,description TEXT NOT NULL,required_skills TEXT NOT NULL DEFAULT '',preferred_skills TEXT NOT NULL DEFAULT '',min_experience REAL NOT NULL DEFAULT 0,status TEXT NOT NULL DEFAULT 'active',created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY,job_id TEXT NOT NULL,user_id TEXT NOT NULL,name TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'queued',total_files INTEGER NOT NULL DEFAULT 0,processed_files INTEGER NOT NULL DEFAULT 0,failed_files INTEGER NOT NULL DEFAULT 0,created_at TEXT NOT NULL,finished_at TEXT,error TEXT);
CREATE TABLE IF NOT EXISTS candidates(id TEXT PRIMARY KEY,batch_id TEXT NOT NULL,job_id TEXT NOT NULL,user_id TEXT NOT NULL,filename TEXT NOT NULL,candidate_name TEXT NOT NULL,resume_text TEXT NOT NULL,skills TEXT NOT NULL,experience_years REAL NOT NULL DEFAULT 0,total_score REAL NOT NULL DEFAULT 0,semantic_score REAL NOT NULL DEFAULT 0,skill_score REAL NOT NULL DEFAULT 0,experience_score REAL NOT NULL DEFAULT 0,matched_skills TEXT NOT NULL DEFAULT '',missing_skills TEXT NOT NULL DEFAULT '',status TEXT NOT NULL DEFAULT 'screened',decision TEXT NOT NULL DEFAULT 'Needs review',created_at TEXT NOT NULL,resume_hash TEXT NOT NULL DEFAULT '',email TEXT NOT NULL DEFAULT '',resume_path TEXT NOT NULL DEFAULT '');
CREATE INDEX IF NOT EXISTS idx_jobs_user ON jobs(user_id);
CREATE INDEX IF NOT EXISTS idx_batches_user ON batches(user_id);
CREATE INDEX IF NOT EXISTS idx_candidates_user ON candidates(user_id);
CREATE INDEX IF NOT EXISTS idx_candidates_batch ON candidates(batch_id);
CREATE INDEX IF NOT EXISTS idx_candidates_job ON candidates(job_id);
"""
with engine.begin() as c:
    for stmt in SCHEMA.split(";"):
        if stmt.strip():
            c.execute(text(stmt))

    # Safely add missing columns for existing SQLite/PostgreSQL databases.
    if engine.dialect.name == "sqlite":
        columns = {
            row[1]
            for row in c.execute(
                text("PRAGMA table_info(candidates)")
            ).fetchall()
        }
    else:
        columns = {
            row[0]
            for row in c.execute(text("""
                SELECT column_name
                FROM information_schema.columns
                WHERE table_schema = current_schema()
                  AND table_name = 'candidates'
            """)).fetchall()
        }

    migrations = [
        ("resume_hash", "VARCHAR(64) NOT NULL DEFAULT ''"),
        ("email", "VARCHAR(320) NOT NULL DEFAULT ''"),
        ("resume_path", "TEXT NOT NULL DEFAULT ''"),
    ]

    for col, ddl in migrations:
        if col not in columns:
            c.execute(
                text(
                    f"ALTER TABLE candidates ADD COLUMN {col} {ddl}"
                )
            )
def now(): return datetime.now(timezone.utc).isoformat()

def query(sql,p=None):

 with engine.begin() as c:

  r=c.execute(text(sql),p or {})

  return [dict(x._mapping) for x in r.fetchall()] if r.returns_rows else []

def execute(sql,p=None):

 with engine.begin() as c: return c.execute(text(sql),p or {}).rowcount

def one(sql,p=None):

 r=query(sql,p); return r[0] if r else None

def get_user(req):

 uid=req.session.get('uid'); return one('SELECT * FROM users WHERE id=:id',{'id':uid}) if uid else None

def need_user(req):

 u=get_user(req)

 if not u: raise HTTPException(401,'Please sign in')

 return u

def render(req,t,ctx=None):

 d={'request':req,'user':get_user(req),'google_configured':GOOGLE}; d.update(ctx or {})

 return templates.TemplateResponse(request=req,name=t,context=d)

def clean(s): return re.sub(r'\s+',' ',s or '').strip()

def description_html(s):

 import html

 t=html.escape(s or '')

 t=re.sub(r'\*\*(.+?)\*\*',r'<strong>\1</strong>',t)

 return '<br>'.join(t.splitlines())

SKILLS = [

 'python','java','javascript','typescript','sql','postgresql','mysql','mongodb',

 'machine learning','deep learning','natural language processing','nlp','spacy',

 'scikit-learn','pandas','numpy','pytorch','tensorflow','fastapi','flask','django',

 'react','html','css','docker','kubernetes','aws','azure','git','linux','rest api',

 'power bi','tableau','excel','data analysis','statistics','data visualization',

 'spark','hadoop','communication','leadership','problem solving','cybersecurity',

 'network security','c++','c#','go','rust'

]

ALIASES={'sklearn':'scikit-learn','postgres':'postgresql','powerbi':'power bi','js':'javascript','ml':'machine learning','dl':'deep learning'}

EXTS={'.pdf','.docx','.txt'}

def extract(filename,data):

 ext=Path(filename).suffix.lower()

 if ext=='.pdf':

  with pymupdf.open(stream=data,filetype='pdf') as pdf: return '\n'.join(page.get_text('text') for page in pdf)

 if ext=='.docx':

  doc=Document(io.BytesIO(data))

  parts=[p.text for p in doc.paragraphs]

  for t in doc.tables:

   for row in t.rows:

     parts.append(' | '.join(c.text for c in row.cells))

  return '\n'.join(parts)

 if ext=='.txt':

   return data.decode('utf-8',errors='ignore')

 raise ValueError('Unsupported format')

def skills_of(s):

 s=clean(s).lower()

 for a,b in ALIASES.items(): s=re.sub(r'(?<!\w)'+re.escape(a)+r'(?!\w)',b,s)

 return sorted({x for x in SKILLS if re.search(r'(?<![a-z0-9+#])'+re.escape(x)+r'(?![a-z0-9+#])',s)})

def experience_of(s):

 for pat in [r'(\d+(?:\.\d+)?)\+?\s+years?\s+(?:of\s+)?(?:professional\s+)?experience',r'experience\s*[:-]\s*(\d+(?:\.\d+)?)\s+years?',r'(\d+(?:\.\d+)?)\+?\s+yrs?\s+experience']:

  m=re.search(pat,s.lower())

  if m:return min(float(m.group(1)),60)

 return 0.0



def name_from_filename(filename):

    """Extract a likely person's name from a resume filename."""



    stem = Path(filename or "").stem

    stem = re.sub(r"^\s*\d{1,3}[.)_\-\s]+", "", stem)

    stem = re.sub(r"[_\-]+", " ", stem)

    stem = re.sub(r"\s+", " ", stem).strip()



    # Stop when the filename reaches a job title or resume descriptor.

    stop_words = {

        "resume", "resumes", "cv", "curriculum", "vitae",

        "professional", "updated", "final", "latest",

        "fresher", "graduate", "entry", "level",

        "years", "year", "experience", "experienced",

        "software", "engineer", "engineering", "developer",

        "analyst", "scientist", "manager", "marketing",

        "designer", "intern", "internship", "backend",

        "frontend", "front", "end", "full", "stack",

        "java", "python", "data", "web", "mobile",

        "cloud", "devops", "security", "network",

        "administrator", "consultant", "specialist",

        "architect", "tester", "testing", "lead",

        "senior", "junior", "principal", "director",

        "officer", "coordinator", "executive", "student",

        "accountant", "sales", "support", "hr",

        "2024", "2025", "2026", "2027", "2028"

    }



    name_parts = []



    for word in stem.split():

        cleaned = re.sub(r"^[^A-Za-z]+|[^A-Za-z'.]+$", "", word)



        if not cleaned:

            break



        if any(char.isdigit() for char in word):

            break



        if cleaned.lower() in stop_words:

            break



        name_parts.append(cleaned)



        # Avoid treating an entire filename as a person's name.

        if len(name_parts) == 4:

            break



    # Require at least two name parts to reduce false positives.

    if len(name_parts) < 2:

        return None



    return " ".join(

        part[0].upper() + part[1:] for part in name_parts

    )





def name_of(s, filename):

    """Prefer a name in the resume; fall back to its filename."""



    role_terms = {

        "software", "engineer", "developer", "manager",

        "analyst", "designer", "consultant", "administrator",

        "architect", "intern", "specialist", "scientist",

        "lead", "director", "marketing", "mechanical",

        "full", "stack", "frontend", "backend", "front-end",

        "back-end", "senior", "junior", "principal",

        "officer", "coordinator", "executive", "student",

        "resume", "curriculum", "vitae"

    }



    for line in [x.strip() for x in s.splitlines()[:10] if x.strip()]:

        line = re.sub(

            r"^\s*(?:resume|curriculum vitae)\s*[:|—-]\s*",

            "",

            line,

            flags=re.I

        )

        line = re.sub(r"^\s*\d{1,3}[.)_\- ]+", "", line)



        if "@" in line or re.search(

            r"\b(?:phone|email|linkedin|github|https?://)\b",

            line,

            re.I

        ):

            continue



        words = re.findall(r"[A-Za-z][A-Za-z.'-]*", line)



        if not words or len(words) > 10:

            continue



        name_words = []



        for word in words:

            if word.lower() in role_terms:

                break

            name_words.append(word)



        candidate = " ".join(name_words).strip()



        if 2 <= len(candidate) <= 45 and 2 <= len(name_words) <= 4:

            return candidate



    # Resume text did not provide a name; try the uploaded filename.

    return (

        name_from_filename(filename)

        or "Candidate (name not detected)"

    )





def score(s,skills,exp,job):

 req={x for x in job['required_skills'].split(';') if x}; pref={x for x in job['preferred_skills'].split(';') if x}; have=set(skills)

 matched=sorted(req&have); missing=sorted(req-have); skill=len(matched)/len(req) if req else 0

 try: mat=TfidfVectorizer(ngram_range=(1,2),stop_words='english').fit_transform([s,job['description']+' '+' '.join(req|pref)]); lexical=float(cosine_similarity(mat[0:1],mat[1:2])[0,0])

 except ValueError: lexical=0

 expfit=1 if float(job['min_experience'])<=0 else min(1,max(0,exp)/float(job['min_experience'])); pref_fit=len(pref&have)/len(pref) if pref else 0

 return {'total':round(100*(.4*skill+.3*lexical+.2*expfit+.1*pref_fit),2),'text':round(lexical*100,2),'skill':round(skill*100,2),'exp':round(expfit*100,2),'matched':matched,'missing':missing}

def zip_entries(data):

 try:z=zipfile.ZipFile(io.BytesIO(data))

 except zipfile.BadZipFile:raise ValueError('Invalid ZIP')

 out=[]; size=0

 with z:

  for info in z.infolist():

   if info.is_dir():continue

   p=PurePosixPath(info.filename.replace('\\\\','/'))

   if p.is_absolute() or '..' in p.parts:raise ValueError('Unsafe ZIP path')

   if Path(p.name).suffix.lower() not in EXTS:continue

   size+=info.file_size

   if size>MAX_ZIP_MB*1024*1024:raise ValueError('ZIP expands beyond configured limit')

   if info.file_size<=MAX_MB*1024*1024:out.append((p.name,z.read(info)))

   if len(out)>MAX_FILES:raise ValueError('Too many files in batch')

 if not out:raise ValueError('No PDF, DOCX, or TXT files found')

 return out

def process_batch(bid, jid, uid, entries):
    batch = one('SELECT * FROM batches WHERE id=:b AND user_id=:u', {'b': bid, 'u': uid})
    job = one('SELECT * FROM jobs WHERE id=:j AND user_id=:u', {'j': jid, 'u': uid})
    if not batch or not job:
        return

    execute("UPDATE batches SET status='processing' WHERE id=:b", {'b': bid})
    good = bad = 0
    seen = set()

    for filename, data in entries:
        stored_file = None
        try:
            safe_name = Path(filename).name
            key = (safe_name.lower(), len(data))
            if key in seen:
                raise ValueError('duplicate')
            seen.add(key)
            raw = clean(extract(safe_name, data))
            if len(raw) < 20:
                raise ValueError('no readable text')

            digest = hashlib.sha256(re.sub(r'\s+', ' ', raw.lower()).encode('utf-8')).hexdigest()
            email_match = re.search(r'[A-Z0-9._%+-]+@[A-Z0-9.-]+.[A-Z]{2,}', raw, re.I)
            email = email_match.group(0).lower() if email_match else ''
            nm = name_of(raw, safe_name)

            existing = one("""SELECT id, resume_path FROM candidates
                WHERE user_id=:u AND job_id=:j AND (
                    resume_hash=:h OR (:email != '' AND lower(email)=:email)
                    OR (:name NOT IN ('Candidate (name not detected)','') AND lower(candidate_name)=lower(:name))
                ) LIMIT 1""", {'u': uid, 'j': jid, 'h': digest, 'email': email, 'name': nm})
            if existing:
                # Backfill original files for duplicate records created before this feature.
                if not existing.get('resume_path'):
                    resume_dir = BASE / 'data' / 'resumes' / uid
                    resume_dir.mkdir(parents=True, exist_ok=True)
                    extension = Path(safe_name).suffix.lower()
                    stored_file = resume_dir / f"{existing['id']}{extension}"
                    stored_file.write_bytes(data)
                    resume_path = stored_file.relative_to(BASE).as_posix()
                    execute('UPDATE candidates SET resume_path=:path WHERE id=:id AND user_id=:u',
                            {'path': resume_path, 'id': existing['id'], 'u': uid})
                    stored_file = None  # Keep the file after successful backfill.
                bad += 1
                execute('UPDATE batches SET processed_files=:p,failed_files=:f WHERE id=:b',
                        {'p': good, 'f': bad, 'b': bid})
                continue

            sk = skills_of(raw)
            exp = experience_of(raw)
            sc = score(raw, sk, exp, job)
            candidate_id = str(uuid.uuid4())
            resume_dir = BASE / 'data' / 'resumes' / uid
            resume_dir.mkdir(parents=True, exist_ok=True)
            extension = Path(safe_name).suffix.lower()
            stored_file = resume_dir / f'{candidate_id}{extension}'
            stored_file.write_bytes(data)
            resume_path = stored_file.relative_to(BASE).as_posix()

            execute("""INSERT INTO candidates(
                id,batch_id,job_id,user_id,filename,candidate_name,resume_text,skills,
                experience_years,total_score,semantic_score,skill_score,experience_score,
                matched_skills,missing_skills,status,decision,created_at,resume_hash,email,resume_path
            ) VALUES (
                :id,:b,:j,:u,:f,:n,:raw,:skills,:exp,:total,:text,:skill,:es,
                :matched,:missing,'screened','Needs review',:created,:hash,:email,:resume_path
            )""", {
                'id': candidate_id, 'b': bid, 'j': jid, 'u': uid,
                'f': safe_name[:255], 'n': nm[:120], 'raw': raw[:100000],
                'skills': ';'.join(sk), 'exp': exp, 'total': sc['total'],
                'text': sc['text'], 'skill': sc['skill'], 'es': sc['exp'],
                'matched': ';'.join(sc['matched']), 'missing': ';'.join(sc['missing']),
                'created': now(), 'hash': digest, 'email': email, 'resume_path': resume_path
            })
            stored_file = None  # File is now associated with the inserted candidate.
            good += 1
        except Exception:
            if stored_file and stored_file.exists():
                try:
                    stored_file.unlink()
                except OSError:
                    pass
            bad += 1

        execute('UPDATE batches SET processed_files=:p,failed_files=:f WHERE id=:b',
                {'p': good, 'f': bad, 'b': bid})

    execute("UPDATE batches SET status='completed',finished_at=:t WHERE id=:b", {'t': now(), 'b': bid})

@app.get('/',response_class=HTMLResponse)

def home(request:Request):

 u=get_user(request)

 if not u:return render(request,'landing.html')

 jobs=query('SELECT * FROM jobs WHERE user_id=:u ORDER BY created_at DESC',{'u':u['id']}); batches=query('SELECT b.*,j.title job_title FROM batches b JOIN jobs j ON j.id=b.job_id WHERE b.user_id=:u ORDER BY b.created_at DESC LIMIT 8',{'u':u['id']})

 counts={'jobs':len(jobs),'candidates':one("SELECT COUNT(*) n FROM (SELECT DISTINCT job_id,COALESCE(NULLIF(email,''),NULLIF(resume_hash,''),lower(candidate_name)) k FROM candidates WHERE user_id=:u)",{'u':u['id']})['n'],'interviews':one("SELECT COUNT(*) n FROM candidates WHERE user_id=:u AND decision='Interview'",{'u':u['id']})['n'],'review':one("SELECT COUNT(*) n FROM candidates WHERE user_id=:u AND decision='Needs review'",{'u':u['id']})['n']}

 return render(request,'dashboard.html',{'jobs':jobs,'batches':batches,'counts':counts})

@app.get('/login',response_class=HTMLResponse)

def login(request:Request):return render(request,'login.html',{'error':request.query_params.get('error')})

@app.get('/auth/google')

async def google_login(request:Request):

 if not GOOGLE:return RedirectResponse('/login?error=google_not_configured',303)

 return await oauth.google.authorize_redirect(request,os.getenv('GOOGLE_REDIRECT_URI') or str(request.url_for('google_callback')),prompt='select_account')

@app.get('/auth/google/callback',name='google_callback')

async def google_callback(request:Request):

 if not GOOGLE:return RedirectResponse('/login?error=google_not_configured',303)

 try:

  token=await oauth.google.authorize_access_token(request); info=token.get('userinfo') or await oauth.google.userinfo(token=token); email=(info.get('email') or '').lower()

  if not info.get('email_verified',False):return RedirectResponse('/login?error=email_not_verified',303)

  if ALLOWED and email not in ALLOWED:return RedirectResponse('/login?error=not_authorized',303)

  old=one('SELECT * FROM users WHERE email=:e',{'e':email}); uid=old['id'] if old else str(uuid.uuid4())

  if old:execute('UPDATE users SET name=:n,picture=:p WHERE id=:id',{'n':info.get('name') or email,'p':info.get('picture') or '','id':uid})

  else:execute('INSERT INTO users(id,email,name,picture,created_at) VALUES(:id,:e,:n,:p,:t)',{'id':uid,'e':email,'n':info.get('name') or email,'p':info.get('picture') or '','t':now()})

  request.session.clear();request.session['uid']=uid;return RedirectResponse('/',303)

 except Exception:return RedirectResponse('/login?error=oauth_failed',303)

@app.post("/auth/demo")
def demo_login(request: Request):
    if not DEMO_LOGIN_ENABLED:
        return RedirectResponse(
            "/login?error=google_not_configured",
            303,
        )

    # One shared account for the public demonstration.
    email = "demo@resumeflow.local"
    name = "Demo Recruiter"

    old = one(
        "SELECT * FROM users WHERE email=:e",
        {"e": email},
    )

    uid = old["id"] if old else str(uuid.uuid4())

    if not old:
        execute(
            """
            INSERT INTO users(id,email,name,picture,created_at)
            VALUES(:id,:e,:n,'',:t)
            """,
            {
                "id": uid,
                "e": email,
                "n": name,
                "t": now(),
            },
        )

    request.session.clear()
    request.session["uid"] = uid
    return RedirectResponse("/", 303)
@app.post('/logout')

def logout(request:Request):request.session.clear();return RedirectResponse('/',303)

@app.get('/jobs/new',response_class=HTMLResponse)

def new_job(request:Request):need_user(request);return render(request,'job_new.html')

@app.post('/jobs')

def create_job(request:Request,title:str=Form(...),description:str=Form(...),required_skills:str=Form(...),preferred_skills:str=Form(''),min_experience:float=Form(0)):

 u=need_user(request);jid=str(uuid.uuid4());req=[x.strip().lower() for x in re.split(r'[,;\n]',required_skills) if x.strip()];pref=[x.strip().lower() for x in re.split(r'[,;\n]',preferred_skills) if x.strip()]

 execute("INSERT INTO jobs(id,user_id,title,description,required_skills,preferred_skills,min_experience,status,created_at) VALUES(:id,:u,:title,:d,:r,:p,:e,'active',:t)",{'id':jid,'u':u['id'],'title':title.strip()[:200],'d':description.strip()[:20000],'r':';'.join(sorted(set(req))),'p':';'.join(sorted(set(pref))),'e':max(0,min(float(min_experience),60)),'t':now()});return RedirectResponse(f'/jobs/{jid}',303)

@app.get('/jobs/{jid}/edit',response_class=HTMLResponse)

def edit_job_page(request:Request,jid:str):

 u=need_user(request); j=one('SELECT * FROM jobs WHERE id=:j AND user_id=:u',{'j':jid,'u':u['id']})

 if not j: raise HTTPException(404,'Campaign not found')

 return render(request,'job_new.html',{'job':j,'editing':True})

@app.post('/jobs/{jid}/edit')

def update_job(request:Request,jid:str,title:str=Form(...),description:str=Form(...),required_skills:str=Form(...),preferred_skills:str=Form(''),min_experience:float=Form(0),status:str=Form('active')):

 u=need_user(request); j=one('SELECT id FROM jobs WHERE id=:j AND user_id=:u',{'j':jid,'u':u['id']})

 if not j: raise HTTPException(404,'Campaign not found')

 if status not in {'active','inactive'}: status='active'

 req=[x.strip().lower() for x in re.split(r'[,;\n]',required_skills) if x.strip()]; pref=[x.strip().lower() for x in re.split(r'[,;\n]',preferred_skills) if x.strip()]

 execute('UPDATE jobs SET title=:title,description=:d,required_skills=:r,preferred_skills=:p,min_experience=:e,status=:s WHERE id=:j AND user_id=:u',{'title':title.strip()[:200],'d':description.strip()[:20000],'r':';'.join(sorted(set(req))),'p':';'.join(sorted(set(pref))),'e':max(0,min(float(min_experience),60)),'s':status,'j':jid,'u':u['id']})

 return RedirectResponse(f'/jobs/{jid}',303)

@app.post('/jobs/{jid}/status')

def update_job_status(request:Request,jid:str,status:str=Form(...)):

 u=need_user(request)

 if status not in {'active','inactive'}: raise HTTPException(400,'Invalid campaign status')

 execute('UPDATE jobs SET status=:s WHERE id=:j AND user_id=:u',{'s':status,'j':jid,'u':u['id']})

 return RedirectResponse('/',303)

@app.get('/jobs/{jid}',response_class=HTMLResponse)

def job_page(request:Request,jid:str):

 u=need_user(request);j=one('SELECT * FROM jobs WHERE id=:j AND user_id=:u',{'j':jid,'u':u['id']})

 if not j:raise HTTPException(404,'Job not found')

 batches=query('SELECT * FROM batches WHERE job_id=:j AND user_id=:u ORDER BY created_at DESC',{'j':jid,'u':u['id']});return render(request,'job_detail.html',{'job':j,'description_display':description_html(j['description']),'batches':batches,'error':request.query_params.get('error'),'max_upload_mb':MAX_MB})

@app.post('/jobs/{jid}/screen')

async def screen(request:Request,jid:str,files:list[UploadFile]=File(...)):

 u=need_user(request);j=one('SELECT * FROM jobs WHERE id=:j AND user_id=:u',{'j':jid,'u':u['id']})

 if not j:raise HTTPException(404,'Job not found')

 if j['status']!='active': return RedirectResponse(f'/jobs/{jid}?error=campaign_inactive',303)

 entries=[]

 for f in files:

  filename=Path(f.filename or 'upload').name;data=await f.read();await f.close()

  if filename.lower().endswith('.zip'):

   try:entries.extend(zip_entries(data))

   except ValueError:return RedirectResponse(f'/jobs/{jid}?error=zip_error',303)

  elif Path(filename).suffix.lower() in EXTS and len(data)<=MAX_MB*1024*1024:entries.append((filename,data))

  if len(entries)>MAX_FILES:return RedirectResponse(f'/jobs/{jid}?error=batch_limit',303)

 if not entries:return RedirectResponse(f'/jobs/{jid}?error=no_supported_files',303)

 bid=str(uuid.uuid4());execute("INSERT INTO batches(id,job_id,user_id,name,status,total_files,processed_files,failed_files,created_at) VALUES(:id,:j,:u,:n,'queued',:total,0,0,:t)",{'id':bid,'j':jid,'u':u['id'],'n':f"{j['title']} · {datetime.now().strftime('%d %b %Y %H:%M')}",'total':len(entries),'t':now()})

 threading.Thread(target=process_batch,args=(bid,jid,u['id'],entries),daemon=True).start();return RedirectResponse(f'/batches/{bid}',303)

@app.get('/batches/{bid}',response_class=HTMLResponse)

def batch_page(request:Request,bid:str):

 u=need_user(request);b=one('SELECT b.*,j.title job_title FROM batches b JOIN jobs j ON j.id=b.job_id WHERE b.id=:b AND b.user_id=:u',{'b':bid,'u':u['id']})

 if not b:raise HTTPException(404,'Batch not found')

 cs=query('SELECT * FROM candidates WHERE batch_id=:b AND user_id=:u ORDER BY total_score DESC',{'b':bid,'u':u['id']});return render(request,'batch.html',{'batch':b,'candidates':cs})

@app.post('/candidates/{cid}/decision')

def decision(request:Request,cid:str,decision:str=Form(...)):

 u=need_user(request)

 if decision not in {'Needs review','Interview','Hold','Not proceeding'}:raise HTTPException(400,'Invalid decision')

 c=one('SELECT * FROM candidates WHERE id=:id AND user_id=:u',{'id':cid,'u':u['id']})

 if not c:raise HTTPException(404,'Candidate not found')

 execute('UPDATE candidates SET decision=:d WHERE id=:id AND user_id=:u',{'d':decision,'id':cid,'u':u['id']});return RedirectResponse(f"/batches/{c['batch_id']}",303)

@app.get('/candidates/{cid}/resume')
def resume_view(request: Request, cid: str):
    u = need_user(request)
    candidate = one('SELECT * FROM candidates WHERE id=:id AND user_id=:u', {'id': cid, 'u': u['id']})
    if not candidate:
        raise HTTPException(404, 'Resume not found')

    relative_path = candidate.get('resume_path') or ''
    if not relative_path:
        raise HTTPException(404, 'Original file unavailable. Upload this resume again to store its original file.')

    resume_root = (BASE / 'data' / 'resumes').resolve()
    file_path = (BASE / relative_path).resolve()
    if not file_path.is_relative_to(resume_root):
        raise HTTPException(403, 'Invalid resume path')
    if not file_path.is_file():
        raise HTTPException(404, 'Original resume file not found')

    filename = Path(candidate.get('filename') or 'resume').name
    if file_path.suffix.lower() == '.pdf':
        return FileResponse(path=str(file_path), media_type='application/pdf', filename=filename,
                            content_disposition_type='inline')

    media_type = {
        '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        '.txt': 'text/plain; charset=utf-8'
    }.get(file_path.suffix.lower(), 'application/octet-stream')
    return FileResponse(path=str(file_path), media_type=media_type, filename=filename,
                        content_disposition_type='attachment')

@app.get('/batches/{bid}/export.xlsx')

def export_xlsx(request:Request,bid:str):

 u=need_user(request); b=one('SELECT * FROM batches WHERE id=:b AND user_id=:u',{'b':bid,'u':u['id']})

 if not b: raise HTTPException(404,'Batch not found')

 rows=query('SELECT candidate_name,filename,total_score,skill_score,experience_score,experience_years,skills,matched_skills,missing_skills,decision,status FROM candidates WHERE batch_id=:b AND user_id=:u ORDER BY total_score DESC',{'b':bid,'u':u['id']})

 fields=['Rank','Applicant name','Resume file','Match score (%)','Skills match (%)','Experience fit (%)','Experience (years)','Detected skills','Matched required skills','Skills to verify','HR decision','Processing status']

 wb=Workbook(); ws=wb.active; ws.title='Screening Results'; ws.append(fields)

 for i,r in enumerate(rows,1): ws.append([i,r['candidate_name'],r['filename'],r['total_score'],r['skill_score'],r['experience_score'],r['experience_years'],r['skills'].replace(';',', '),r['matched_skills'].replace(';',', '),r['missing_skills'].replace(';',', '),r['decision'],r['status']])

 for cell in ws[1]: cell.font=Font(bold=True,color='242321'); cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cell.fill=PatternFill('solid',fgColor='E9E2D6')

 for row in ws.iter_rows(min_row=2):

  for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)

 for i,w in enumerate([8,25,34,17,17,18,18,38,35,35,20,20],1): ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=w

 ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=32

 out=io.BytesIO(); wb.save(out); out.seek(0)

 return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="resumeflow-{bid[:8]}.xlsx"'})



@app.get('/batches/{bid}/export.csv')

def export(request:Request,bid:str):

 u=need_user(request);b=one('SELECT * FROM batches WHERE id=:b AND user_id=:u',{'b':bid,'u':u['id']})

 if not b:raise HTTPException(404,'Batch not found')

 rows=query('SELECT candidate_name,filename,total_score,skill_score,semantic_score,experience_score,experience_years,skills,matched_skills,missing_skills,decision,status FROM candidates WHERE batch_id=:b AND user_id=:u ORDER BY total_score DESC',{'b':bid,'u':u['id']});out=io.StringIO();fields=list(rows[0]) if rows else ['candidate_name','filename','total_score','skill_score','semantic_score','experience_score','experience_years','skills','matched_skills','missing_skills','decision','status'];w=csv.DictWriter(out,fieldnames=fields);w.writeheader();w.writerows(rows);out.seek(0)

 return StreamingResponse(iter([out.getvalue()]),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="resumeflow-{bid[:8]}.csv"'})

@app.get('/api/batches/{bid}')

def batch_api(request:Request,bid:str):

 u=need_user(request);b=one('SELECT * FROM batches WHERE id=:b AND user_id=:u',{'b':bid,'u':u['id']})

 if not b:raise HTTPException(404,'Batch not found')

 return {k:b[k] for k in ['id','status','total_files','processed_files','failed_files','finished_at','error']}

@app.get('/health')

def health():

 try:

  with engine.connect() as c:c.execute(text('SELECT 1'))

  return {'status':'ok','database':'connected','service':'ResumeFlow'}

 except Exception:return JSONResponse({'status':'error','database':'unavailable'},status_code=503)
